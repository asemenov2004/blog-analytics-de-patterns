"""Сборка отчёта .docx по лабораторной (по образцу отчётов VBA/Лаб8/automation/make_report.py).

Оформление: титульник ЮФУ из template_title.docx, поля 3 / 1,5 / 2 / 2 см, Times New Roman 14,
подписи «Рисунок N – …», «Таблица N – …», «Листинг N – …».

    from docx_report import Report
    R = Report(lab=1, topic="Паттерны приёма данных", out=Path("Отчет_Лаб1_ТКР_Семенов.docx"))
    R.para("Цель работы: ...")
    R.save()
"""
from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.shared import Cm, Pt
from PIL import Image

from meta import fix_meta

LIB = Path(__file__).resolve().parent
TEMPLATE = LIB / "template_title.docx"
DISCIPLINE = "Технологии коллективной разработки"
# Преподаватель: (должность с «Проверил/Проверила», ФИО). None — оставить пустую строку под подпись.
# TODO: уточнить должность и форму «Проверил/Проверила», напр. ("Проверил доцент", ...).
TEACHER: tuple[str, str] | None = ("Проверил(а)", "Чернышев-Придворов Ю. Г.")


def set_text(p, text: str) -> None:
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


class Report:
    def __init__(self, lab: int, topic: str, out: Path, teacher: tuple[str, str] | None = TEACHER):
        self.d = docx.Document(TEMPLATE)
        self.out = out
        self.fig = self.tab = self.lst = 0
        self._title_page(lab, topic, teacher)

    def _title_page(self, lab: int, topic: str, teacher) -> None:
        for p in self.d.paragraphs:
            t = p.text
            if t.startswith("Теория информационных"):
                set_text(p, DISCIPLINE)
            elif t.startswith("Отчёт по лабораторной работе"):
                set_text(p, f"Отчёт по лабораторной работе № {lab}")
            elif t.startswith("«Исследование"):
                set_text(p, f"«{topic}»")
            elif t.startswith("Вариант"):
                set_text(p, "")
            elif t.startswith("Проверил"):
                who, name = teacher or ("Проверил", "")
                # отступ как у строки студента, но длинная фамилия «съедает» часть пробелов, иначе строка
                # переносится (шрифт пропорциональный — выравнивание приблизительное)
                student = next(x.text for x in self.d.paragraphs if x.text.startswith("Выполнил"))
                k, student_name = student.index("_"), student.rsplit("     ", 1)[-1]
                pad = max(2, k - len(who) - max(0, len(name) - len(student_name)))
                set_text(p, who + " " * pad + "_______________     " + name)

    # --- текст
    def para(self, text: str = "", align=AL.JUSTIFY, bold=False, size=None, indent=True, italic=False,
             after=0, keep=False):
        p = self.d.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        if not indent:
            pf.first_line_indent = Cm(0)
        pf.space_after = Pt(after)
        pf.keep_with_next = keep
        if text:
            self.runs(p, text, bold, size, italic)
        return p

    @staticmethod
    def runs(p, text: str, bold=False, size=None, italic=False) -> None:
        """`код` в тексте выводится моноширинным шрифтом."""
        for i, part in enumerate(re.split(r"`", text)):
            if not part:
                continue
            r = p.add_run(part)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
            if i % 2:
                r.font.name = "Courier New"
                r.font.size = Pt((size or 14) - 2)

    def h1(self, text: str) -> None:
        self.para(text, AL.CENTER, bold=True, indent=False, keep=True, after=6)

    def h2(self, text: str) -> None:
        self.para(text, AL.LEFT, bold=True, indent=False, keep=True, after=3)

    def bullets(self, items: list[str]) -> None:
        for it in items:
            p = self.d.add_paragraph(style="List Bullet")
            p.alignment = AL.JUSTIFY
            self.runs(p, it)

    # --- нумерация для ссылок «(рисунок N)» до вставки объекта
    def next_fig(self, k: int = 1) -> int:
        return self.fig + k

    def next_tab(self, k: int = 1) -> int:
        return self.tab + k

    def next_lst(self, k: int = 1) -> int:
        return self.lst + k

    # --- объекты
    def figure(self, path: Path, caption: str, max_cm: float = 16.0, scale: float = 0.85) -> int:
        with Image.open(path) as im:
            w_cm = min(max_cm, im.width * 2.54 / 96 * scale)
        p = self.para(align=AL.CENTER, indent=False, keep=True)
        p.add_run().add_picture(str(path), width=Cm(w_cm))
        self.fig += 1
        self.para(f"Рисунок {self.fig} – {caption}", AL.CENTER, size=12, indent=False, after=8)
        return self.fig

    def table(self, rows: list[list], caption: str, widths: list[float] | None = None, size: int = 11) -> int:
        self.tab += 1
        self.para(f"Таблица {self.tab} – {caption}", AL.LEFT, size=12, indent=False, keep=True, after=2)
        ncol = max(len(r) for r in rows)
        t = self.d.add_table(rows=len(rows), cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j in range(ncol):
                cell = t.cell(i, j)
                p = cell.paragraphs[0]
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = i < len(rows) - 1
                p.alignment = AL.LEFT if j == 0 and i > 0 else AL.CENTER
                r = p.add_run(str(row[j]) if j < len(row) else "")
                r.font.size = Pt(size)
                r.bold = i == 0
                if widths:
                    cell.width = Cm(widths[j])
        self.para(after=4)
        return self.tab

    def listing(self, code: str, caption: str) -> int:
        self.lst += 1
        self.para(f"Листинг {self.lst} – {caption}", AL.LEFT, size=12, indent=False, keep=True, after=2)
        lines = code.rstrip().split("\n")
        for i, line in enumerate(lines):
            p = self.para(align=AL.LEFT, indent=False, keep=i < 3 or i >= len(lines) - 3)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(line if line.strip() else " ")
            r.font.name = "Courier New"
            r.font.size = Pt(9)
        self.para(after=4)
        return self.lst

    def save(self) -> Path:
        self.d.save(self.out)
        fix_meta(self.out)
        return self.out
